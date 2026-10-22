import csv
import json
import math
import os
import re
import sys

# Build site_data.js (window.EXTRA_SITE_DATA) for one community from the analysis outputs.
#   python3 build_site_data.py communities/<id>/build_config.json
# Everything specific to a community (input paths, candidate MAGs, tree choices, stations
# and how sample names map to them) comes from build_config.json; see docs/workflow.md.
CFG_PATH = os.path.abspath(sys.argv[1] if len(sys.argv) > 1 else "build_config.json")
CFG = json.load(open(CFG_PATH, encoding="utf-8"))
def cfg_path(p):
    """config paths: ~ and $VARS expanded, relative paths are relative to the config file"""
    p = os.path.expanduser(os.path.expandvars(p))
    return p if os.path.isabs(p) else os.path.join(os.path.dirname(CFG_PATH), p)

scratch = cfg_path(CFG["inputs_dir"])
out_js = cfg_path(CFG.get("out_js", "site_data.js"))
LOCAL = CFG["local"]            # the community's own sampling: station name, sites, labels
LOCAL_STATION = LOCAL["station"]

mar_cands = CFG["candidates"].get("marine", [])
fw_cands = CFG["candidates"].get("freshwater", [])
# Stand-alone genomes (e.g. reassemblies, community-only bins) mapped and annotated on their
# own: scripts/hpc/genome_site_inputs.sh puts their per-MAG files in one folder each.
EXTRA = {g: dict(v, dir=cfg_path(v["dir"])) for g, v in CFG.get("extra_genomes", {}).items()}
for g, v in EXTRA.items():
    lst = mar_cands if v["habitat"] == "marine" else fw_cands
    if g not in lst:
        lst.append(g)


# optional KofamKOALA ko_list: KO definitions for genes whose annotation has no description
KO_NAMES = {}
if CFG.get("ko_list") and os.path.exists(cfg_path(CFG["ko_list"])):
    for r in csv.DictReader(open(cfg_path(CFG["ko_list"])), delimiter="\t"):
        KO_NAMES[r["knum"]] = r["definition"]

def mag_dir(hab, m):
    return EXTRA[m]["dir"] if m in EXTRA else f"{scratch}/{hab}"

# ---------------------------------------------------------
# 1. PARSE NEWICK & EXTRACT EXACT GTDB SUBTREE
# ---------------------------------------------------------
class Node:
    def __init__(self, name="", length=0.0, children=None):
        self.name = name
        self.length = length
        self.children = children or []
        self.parent = None

def parse_newick(s):
    s = s.strip().rstrip(";")
    pos = 0
    n = len(s)

    def _read_label():
        nonlocal pos
        if pos < n and s[pos] == "'":
            pos += 1
            buf = []
            while pos < n and s[pos] != "'":
                buf.append(s[pos])
                pos += 1
            if pos < n and s[pos] == "'":
                pos += 1
            return "".join(buf).strip()
        else:
            buf = []
            while pos < n and s[pos] not in ",():;":
                buf.append(s[pos])
                pos += 1
            return "".join(buf).strip()

    def _parse():
        nonlocal pos
        children = []
        if pos < n and s[pos] == "(":
            pos += 1
            while True:
                child = _parse()
                children.append(child)
                if pos < n and s[pos] == ",":
                    pos += 1
                elif pos < n and s[pos] == ")":
                    pos += 1
                    break
                else:
                    break
        name = _read_label()
        length = 0.0
        if pos < n and s[pos] == ":":
            pos += 1
            lstr = _read_label()
            try:
                length = float(lstr)
            except:
                length = 0.0
        node = Node(name=name, length=length, children=children)
        for c in children:
            c.parent = node
        return node

    return _parse()

def get_leaves(node):
    if not node.children:
        return [node]
    res = []
    for c in node.children:
        res.extend(get_leaves(c))
    return res

def find_target_leaf(root, queries):
    leaves = get_leaves(root)
    for q in queries:
        if not q: continue
        for lf in leaves:
            if lf.name in (q, f"GB_{q}", f"RS_{q}"):
                return lf
    return None

def prune_node_to_max_leaves(node, target_leaf, max_leaves=14):
    leaves = get_leaves(node)
    if len(leaves) <= max_leaves:
        return node
    new_children = []
    for c in node.children:
        c_leaves = get_leaves(c)
        if target_leaf in c_leaves:
            new_children.append(prune_node_to_max_leaves(c, target_leaf, max(6, max_leaves - len(node.children) + 1)))
        else:
            rep_leaves = c_leaves[:3]
            if len(rep_leaves) == 1:
                new_children.append(rep_leaves[0])
            else:
                sub = Node(name=c.name, length=c.length, children=rep_leaves)
                new_children.append(sub)
    return Node(name=node.name, length=node.length, children=new_children)

def extract_clade(root, queries, min_leaves=9, max_leaves=14):
    leaf = find_target_leaf(root, queries)
    if not leaf:
        return None, None
    curr = leaf
    while curr.parent is not None:
        p = curr.parent
        leaves = get_leaves(p)
        if len(leaves) >= min_leaves:
            curr = p
            break
        curr = p
    curr = prune_node_to_max_leaves(curr, leaf, max_leaves=max_leaves)
    return curr, leaf

import sys
sys.setrecursionlimit(20000)

# GTDB r226 species representatives: taxonomy and type-strain status
gtdb_reps = {}
with open(f"{scratch}/gtdb_r226_reps.tsv") as f:
    for r in csv.DictReader(f, delimiter="\t"):
        sp = r["gtdb_taxonomy"].split(";s__")[-1]
        # a named type strain needs NCBI type material AND a Latin binomial in GTDB
        # (NCBI flags a few MAGs with placeholder GTDB names as type material by mistake)
        named = re.fullmatch(r"[A-Z][a-z]+(_[A-Z]+)? [a-z]+(_[A-Z]+)?", sp) is not None
        typ = r["gtdb_type_designation_ncbi_taxa"] if named else "not type material"
        gtdb_reps[r["accession"]] = (r["gtdb_taxonomy"], typ, r["gtdb_type_species_of_genus"] == "t")

# optional `gtdb_taxonomy`: the full taxonomy file of the GTDB release used for the newer genomes
# (GTDB-Tk db taxonomy/gtdb_taxonomy.tsv), to name tree leaves that the reps table lacks
gtdb_tax_all = {}
if CFG.get("gtdb_taxonomy"):
    for line in open(cfg_path(CFG["gtdb_taxonomy"])):
        a, t = line.rstrip("\n").split("\t")[:2]
        gtdb_tax_all[a] = t

# optional `ani_table`: skani search of every site genome against the GTDB species reps
# (skani search -d <gtdbtk_db>/skani/database ...). Gives the closest relative and the ANI
# shown next to each reference on the tree, so the two always agree.
ANI_HITS = {}
if CFG.get("ani_table"):
    for r in csv.DictReader(open(cfg_path(CFG["ani_table"])), delimiter="\t"):
        q = os.path.basename(r["Query_file"]).rsplit(".", 1)[0]
        acc = re.search(r"(GC[AF]_\d+\.\d+)", os.path.basename(r["Ref_file"])).group(1)
        tax = gtdb_tax_all.get("RS_" + acc) or gtdb_tax_all.get("GB_" + acc) or ""
        ANI_HITS.setdefault(q, []).append({"acc": acc, "species": tax.split(";s__")[-1] if tax else acc,
                                           "ani": float(r["ANI"]), "af": round(float(r["Align_fraction_query"]) / 100, 3)})
    for h in ANI_HITS.values():
        h.sort(key=lambda x: -x["ani"])


def tree_ani(d, hits):
    """Write each reference leaf's ANI to our genome (from ANI_HITS) into the tree dict.
    Leaves are matched by accession or by species name; a species code such as sp002470565
    survives genus renames between GTDB releases, so it also matches."""
    if d.get("children"):
        for c in d["children"]:
            tree_ani(c, hits)
        return
    if d.get("is_target") or d.get("is_user_bin"):
        return
    acc = re.sub(r"^(GB|RS)_", "", d.get("accession") or d.get("name", ""))
    sp = d.get("label", "").split(" [")[0]
    code = re.search(r"\bsp\d{9}\b", sp)
    for h in hits:
        if h["acc"] == acc or h["species"] == sp or (code and h["species"].endswith(" " + code.group(0))):
            d["ani"], d["af"] = h["ani"], h["af"]
            return


def is_ref(name):
    return name.startswith(("GB_GC", "RS_GC"))


def leaf_distances(target):
    """Patristic distance from the target leaf to every leaf (walks the tree as a graph)."""
    dist, stack, out = {id(target): 0.0}, [target], {}
    while stack:
        n = stack.pop()
        d = dist[id(n)]
        if not n.children:
            out[n] = d
        nbrs = [(c, c.length) for c in n.children] + ([(n.parent, n.length)] if n.parent else [])
        for m, l in nbrs:
            if id(m) not in dist:
                dist[id(m)] = d + (l or 0.0)
                stack.append(m)
    return out


def select_tree(root, target, n_refs=8, n_type=2, n_user=3):
    """Prune the GTDB-Tk classify tree to: the target, its n_refs nearest GTDB references,
    the n_type nearest type-strain genomes of named species, and up to n_user nearby
    local bins closer than the 5th-nearest reference."""
    d = leaf_distances(target)
    refs = sorted((l for l in d if is_ref(l.name) and l is not target), key=d.get)
    keep = {target} | set(refs[:n_refs])
    typed = [l for l in refs if gtdb_reps.get(l.name, ("", ""))[1] == "type strain of species"]
    keep |= set(typed[:n_type])
    cutoff = d[refs[min(4, len(refs) - 1)]] if refs else 0
    # other local bins nearby, skipping other binners' copies of the target genome
    base = target.name.replace("dastool-", "")
    users = sorted((l for l in d if not is_ref(l.name) and l is not target and 0.02 <= d[l] <= cutoff
                    and l.name.replace("dastool-", "") != base), key=d.get)
    keep |= set(users[:n_user])

    def induce(n):
        if not n.children:
            return Node(n.name, n.length) if n in keep else None
        kids = [k for k in (induce(c) for c in n.children) if k]
        if not kids:
            return None
        if len(kids) == 1:  # collapse unary nodes, carrying the branch length down
            kids[0].length += n.length
            return kids[0]
        m = Node(n.name, n.length, kids)
        for k in kids:
            k.parent = m
        return m
    sub = induce(root)
    sub.length = 0.0
    new_target = next(l for l in get_leaves(sub) if l.name == target.name)
    return sub, new_target, len(typed[:n_type])


def gtdb_label(acc):
    tax, typ, tsg = gtdb_reps.get(acc) or (gtdb_tax_all.get(acc, ""), "", False)
    sp = tax.split(";")[-1].replace("s__", "") if tax else acc
    short = acc.replace("GB_", "").replace("RS_", "")
    return sp, short, typ == "type strain of species", tax


def node_to_dict(node, target_leaf_name, leaf_meta_map, ani_sister=None):
    if not node.children:
        raw_name = node.name
        info = leaf_meta_map.get(raw_name, {})
        if is_ref(raw_name) and (raw_name in gtdb_reps or raw_name in gtdb_tax_all):
            sp, short, typ, tax = gtdb_label(raw_name)
            info = {"label": f"{sp} [{short}]", "accession": short, "lineage": tax, "is_user_bin": False, "is_type": typ}
        is_target = (raw_name == target_leaf_name) and (ani_sister is None)
        d = {
            "name": raw_name,
            "label": info.get("label", raw_name),
            "accession": info.get("accession", ""),
            "lineage": info.get("lineage", ""),
            "is_user_bin": info.get("is_user_bin", False),
            "is_type": info.get("is_type", False),
            "is_target": is_target,
            "length": round(node.length, 5),
        }
        if raw_name == target_leaf_name and ani_sister is not None:
            dist = round(max(0.002, (100.0 - ani_sister["ani"]) / 100.0), 5)
            return {
                "name": "100.0",
                "support": "100",
                "taxon_tag": f"ANI {ani_sister['ani']}%",
                "length": round(max(0.003, node.length * 0.8), 5),
                "children": [
                    {
                        "name": ani_sister["mag_id"],
                        "label": f"★ {ani_sister['mag_id']} ({ani_sister['nickname']})",
                        "accession": f"{LOCAL['label']} MAG · ANI {ani_sister['ani']}%",
                        "lineage": ani_sister["lineage"],
                        "is_user_bin": True,
                        "is_target": True,
                        "length": dist,
                    },
                    {
                        "name": raw_name,
                        "label": info.get("label", raw_name),
                        "accession": info.get("accession", raw_name),
                        "lineage": info.get("lineage", ""),
                        "is_user_bin": False,
                        "is_type": info.get("is_type", False),
                        "is_target": False,
                        "length": round(max(0.002, node.length * 0.2), 5),
                    }
                ]
            }
        if is_target:
            d["label"] = f"★ {raw_name} ({LOCAL['novel_label']})"
            d["accession"] = "De novo GTDB-Tk placement"
            d["is_user_bin"] = True
        return d
    else:
        sup = node.name.split(":")[0] if node.name else ""
        taxon_tag = node.name.split(":", 1)[1] if ":" in (node.name or "") else ""
        return {
            "name": node.name,
            "support": sup,
            "taxon_tag": taxon_tag,
            "length": round(node.length, 5),
            "children": [node_to_dict(c, target_leaf_name, leaf_meta_map, ani_sister) for c in node.children]
        }

with open(f"{scratch}/marine/phylotree.json") as f:
    mpt = json.load(f)

marine_leaf_meta = {}
for lname, linfo in mpt.get("leaf_info", {}).items():
    acc_raw = linfo.get("id", "")
    acc_clean = acc_raw.replace("GB_", "").replace("RS_", "")
    is_ub = bool(linfo.get("user_bin", False))
    lin = linfo.get("lineage", "")
    sp_part = lin.split(";")[-1].replace("s__", "") if ";" in lin else ""
    genus_part = lin.split(";")[-2].replace("g__", "") if len(lin.split(";")) >= 2 else ""
    if is_ub:
        lbl = f"{lname} ({sp_part or genus_part or CFG.get("coassembly_label", {}).get("marine", "Marine MAG")})"
    else:
        lbl = f"{lname} [{acc_clean}]" if acc_clean and acc_clean not in lname else lname
    entry = {
        "label": lbl,
        "accession": CFG.get("coassembly_label", {}).get("marine_accession", "Marine co-assembly MAG") if is_ub else acc_clean,
        "lineage": lin,
        "is_user_bin": is_ub,
    }
    marine_leaf_meta[lname] = entry
    if acc_raw:
        marine_leaf_meta[acc_raw] = entry
    if acc_clean:
        marine_leaf_meta[acc_clean] = entry

fw_leaf_meta = dict(marine_leaf_meta)
with open(f"{scratch}/freshwater/gtdbtk_taxonomy.tsv") as f:
    for r in csv.DictReader(f, delimiter="\t"):
        ug = r["user_genome"]
        short = ug.replace("dastool-", "")
        cls = r["classification"]
        sp = cls.split(";")[-1].replace("s__", "") or (cls.split(";")[-2].replace("g__", "") + " sp. nov.")
        for k in [ug, short]:
            fw_leaf_meta[k] = {
                "label": f"{ug} ({sp})",
                "accession": "Arctic Freshwater Co-Assembly MAG",
                "lineage": cls,
                "is_user_bin": True,
            }
        for ref_col, tax_col in [("closest_placement_reference", "closest_placement_taxonomy"), ("closest_genome_reference", "closest_genome_taxonomy")]:
            cpr = r.get(ref_col, "")
            cpt = r.get(tax_col, "")
            if cpr and cpr != "N/A":
                sp_ref = cpt.split(";")[-1].replace("s__", "")
                for k in [cpr, f"GB_{cpr}", f"RS_{cpr}"]:
                    fw_leaf_meta[k] = {
                        "label": f"{sp_ref} [{cpr}]" if sp_ref else cpr,
                        "accession": cpr,
                        "lineage": cpt,
                        "is_user_bin": False,
                    }
        others = r.get("other_related_references(genome_id,species_name,radius,ANI,AF)", "")
        if others and others != "N/A":
            for item in others.split(";"):
                parts = [x.strip() for x in item.split(",")]
                if len(parts) >= 2:
                    gid, sname = parts[0], parts[1].replace("s__", "")
                    for k in [gid, f"GB_{gid}", f"RS_{gid}"]:
                        if k not in fw_leaf_meta:
                            fw_leaf_meta[k] = {
                                "label": f"{sname} [{gid}]",
                                "accession": gid,
                                "lineage": "",
                                "is_user_bin": False,
                            }

# optional `gtdbtk_runs`: folders of newer GTDB-Tk classify_wf runs, each holding
# gtdbtk.bac120.summary.tsv, gtdbtk.bac120.tree.mapping.tsv and the classify trees side by side.
# A site genome found in one takes its classification and its tree from that run instead of
# its habitat's original run (e.g. to move every microbe to the same GTDB release).
GTDBTK_RUNS = {}
for run_dir in CFG.get("gtdbtk_runs", []):
    run_dir = cfg_path(run_dir)
    run_trees = {}
    if os.path.exists(f"{run_dir}/gtdbtk.bac120.tree.mapping.tsv"):
        for r in csv.DictReader(open(f"{run_dir}/gtdbtk.bac120.tree.mapping.tsv"), delimiter="\t"):
            run_trees[r.get("User_genome") or r["user_genome"]] = f"{run_dir}/gtdbtk.bac120.classify.tree.{r['class_tree_mapped']}.tree"
    for r in csv.DictReader(open(f"{run_dir}/gtdbtk.bac120.summary.tsv"), delimiter="\t"):
        GTDBTK_RUNS[r["user_genome"]] = (r, run_trees.get(r["user_genome"]))

# per MAG: which GTDB-Tk classify tree to cut, which leaf names to look for, and (for MAGs
# placed by ANI only) the sister species to draw next to it
tree_specs = {m: (v["habitat"], v["tree_file"], v["queries"], v.get("ani_sister"))
              for m, v in CFG["tree_specs"].items()}
for g, v in EXTRA.items():  # the GTDB-Tk classify tree its own run placed it in
    if g not in tree_specs and v.get("tree"):
        tree_specs[g] = (v["habitat"], cfg_path(v["tree"]), [g], None)
hab_of = {m: h for h, ms in CFG["candidates"].items() for m in ms}
for g, (r, tree) in GTDBTK_RUNS.items():  # newer runs override (placed in the tree, no ANI sister)
    if g in hab_of and tree:
        tree_specs[g] = (hab_of[g], tree, [g], None)

extracted_trees = {}
for mag_id, (hab, tfile, queries, ani_sis) in tree_specs.items():
    with open(tfile if os.path.isabs(tfile) else f"{scratch}/{hab}/gtdbtk_trees/{tfile}") as f:
        nw_str = f.read()
    tfile = os.path.basename(tfile)
    lmeta = marine_leaf_meta if hab == "marine" else fw_leaf_meta
    root = parse_newick(nw_str)
    tleaf0 = find_target_leaf(root, queries)
    clade, tleaf, n_typed = select_tree(root, tleaf0) if tleaf0 else (None, None, 0)
    if clade and tleaf:
        tdict = node_to_dict(clade, tleaf.name, lmeta, ani_sister=ani_sis)
        if mag_id in ANI_HITS:
            tree_ani(tdict, ANI_HITS[mag_id])
        n_leaves = len(get_leaves(clade)) + (1 if ani_sis else 0)
        extracted_trees[mag_id] = {
            "tree_file": tfile,
            "n_leaves": n_leaves,
            "placement_type": f"GTDB-Tk r226 ({tfile}) · ANI Sister Placement ({ani_sis['ani']}%)" if ani_sis else f"GTDB-Tk r226 ({tfile}) · De Novo Phylogenetic Placement (pplacer)",
            "n_type_strains": n_typed,
            "root": tdict,
        }
        print(f"Tree {mag_id}: extracted {n_leaves} leaves around '{tleaf.name}' from {tfile}")

# ---------------------------------------------------------
# 2. CIRCOS GENOME TRACKS
# ---------------------------------------------------------
rrna_by_contig = {"marine": {}, "freshwater": {}}
trna_by_contig = {"marine": {}, "freshwater": {}}
bgc_by_contig  = {"marine": {}, "freshwater": {}}

for hab in ["marine", "freshwater"]:
    with open(f"{scratch}/{hab}/rrna_genes.tsv") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            # barrnap runs with all three kingdom models; keep the bacterial calls only
            if r.get("best") == "False" or r.get("kingdom", "Bacteria") != "Bacteria": continue
            rrna_by_contig[hab].setdefault(r["contig_id"], []).append(r)
    with open(f"{scratch}/{hab}/trna_genes.tsv") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            trna_by_contig[hab].setdefault(r["contig_id"], []).append(r)
    with open(f"{scratch}/{hab}/antismash_summary.tsv") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            bgc_by_contig[hab].setdefault(r["contig"], []).append(r)

# dnaA (chromosomal replication initiator) positions from the co-assembly Bakta annotation;
# Bakta also tags the Hda regulator as "dnaA", so filter on the product.
dnaa_by_contig = {"marine": {}, "freshwater": {}}
for hab in ["marine", "freshwater"]:
    p = f"{scratch}/{hab}/dnaA.tsv"
    if os.path.exists(p):
        for r in csv.reader(open(p), delimiter="\t"):
            if len(r) >= 8 and "replication initiat" in r[7]:
                dnaa_by_contig[hab][r[0]] = int(r[2])
for g, v in EXTRA.items():  # their contigs are prefixed with the genome id, so no clashes
    hab, d = v["habitat"], v["dir"]
    domain = v.get("domain", "Bacteria")  # barrnap kingdom to keep; "Archaea" for an archaeon
    for r in csv.DictReader(open(f"{d}/rrna_genes.tsv"), delimiter="\t"):
        if r.get("best") != "False" and r.get("kingdom", "Bacteria") == domain:
            rrna_by_contig[hab].setdefault(r["contig_id"], []).append(r)
    for r in csv.DictReader(open(f"{d}/trna_genes.tsv"), delimiter="\t"):
        trna_by_contig[hab].setdefault(r["contig_id"], []).append(r)
    for r in csv.DictReader(open(f"{d}/antismash_summary.tsv"), delimiter="\t"):
        if r.get("type") not in (None, "", "NA"):
            bgc_by_contig[hab].setdefault(r["contig"], []).append(r)
    for r in csv.reader(open(f"{d}/dnaA.tsv"), delimiter="\t"):
        if len(r) >= 8 and "replication initiat" in r[7]:
            dnaa_by_contig[hab][r[0]] = int(r[2])


def cumulative_skew(seq, win=2000):
    """Cumulative GC skew at window resolution: list of (position, cumulative skew)."""
    out, cum = [(0, 0.0)], 0.0
    for i in range(0, len(seq), win):
        w = seq[i:i + win]
        g, c = w.count("G"), w.count("C")
        cum += (g - c) / max(1, g + c)
        out.append((min(i + win, len(seq)), cum))
    return out


def order_by_skew(contigs, dnaa):
    """Order contigs around a circle from GC skew, ori first.

    The ori sits at the minimum of cumulative GC skew (leading strand G>C after it),
    which holds whichever way a contig is oriented. The ori contig goes first; contigs
    with net positive skew follow (right replichore), then the contig holding the
    terminus (cumulative-skew maximum), then net-negative contigs (left replichore).
    Order within each replichore isn't recoverable from skew, so it is by length.
    Returns (ordered contigs, ori position within the first contig or None, method).
    """
    info = {}
    for cid, seq in contigs:
        cs = cumulative_skew(seq)
        vals = [v for _, v in cs]
        imin, imax = vals.index(min(vals)), vals.index(max(vals))
        edge = max(1, len(vals) // 20)
        g, c = seq.count("G"), seq.count("C")
        info[cid] = dict(cs=cs, net=(g - c) / max(1, g + c),
                         min_depth=min(vals[0] - vals[imin], vals[-1] - vals[imin]) if edge <= imin < len(vals) - edge else 0,
                         max_depth=min(vals[imax] - vals[0], vals[imax] - vals[-1]) if edge <= imax < len(vals) - edge else 0,
                         min_pos=cs[imin][0])
    ori_cid, ori_pos, method = None, None, "none"
    dn = [cid for cid, _ in contigs if cid in dnaa]
    if dn:
        ori_cid = max(dn, key=lambda x: len(dict(contigs)[x]))
        d = dnaa[ori_cid]
        near = info[ori_cid]["min_depth"] > 0 and abs(info[ori_cid]["min_pos"] - d) <= 100_000
        ori_pos, method = (info[ori_cid]["min_pos"], "dnaA + GC skew minimum") if near else (d, "dnaA")
    else:
        best = max(info, key=lambda x: info[x]["min_depth"])
        if info[best]["min_depth"] >= 3:  # a clear V in cumulative skew (~3 windows' worth of full skew)
            ori_cid, ori_pos, method = best, info[best]["min_pos"], "GC skew minimum"
    rest = [(cid, seq) for cid, seq in contigs if cid != ori_cid]
    ter = max(rest, key=lambda x: info[x[0]]["max_depth"], default=None)
    ter = ter if ter and info[ter[0]]["max_depth"] >= 3 else None
    right = sorted((x for x in rest if x is not ter and info[x[0]]["net"] >= 0), key=lambda x: -len(x[1]))
    left = sorted((x for x in rest if x is not ter and info[x[0]]["net"] < 0), key=lambda x: -len(x[1]))
    head = [x for x in contigs if x[0] == ori_cid]
    return head + right + ([ter] if ter else []) + left, ori_pos, method


def compute_circos(hab, mag_id, n_bins=120):
    fa_path = f"{mag_dir(hab, mag_id)}/{mag_id}.fa"
    tsv_path = f"{mag_dir(hab, mag_id)}/{mag_id}.tsv"

    contigs = []
    curr_id = None
    curr_seq = []
    with open(fa_path) as f:
        for line in f:
            line = line.strip()
            if not line: continue
            if line.startswith(">"):
                if curr_id is not None:
                    contigs.append((curr_id, "".join(curr_seq).upper()))
                curr_id = line[1:].split()[0]
                curr_seq = []
            else:
                curr_seq.append(line)
        if curr_id is not None:
            contigs.append((curr_id, "".join(curr_seq).upper()))

    contigs.sort(key=lambda x: -len(x[1]))
    contigs, ori_pos, ori_method = order_by_skew(contigs, dnaa_by_contig[hab])
    total_len = sum(len(s) for _, s in contigs)
    rot = ori_pos or 0  # rotate so the ori is at the top (position 0)

    contig_meta = []
    contig_offsets = {}
    contig_lengths = {}
    offset = 0
    full_seq_parts = []
    for cid, seq in contigs:
        clen = len(seq)
        g = seq.count("G")
        c = seq.count("C")
        gc = round((g + c) / max(1, clen) * 100.0, 1)
        contig_meta.append({
            "id": cid,
            "length": clen,
            "start": offset - rot,
            "end": offset + clen - rot,
            "gc": gc,
        })
        contig_offsets[cid] = offset - rot  # may be negative for the ori contig; positions wrap below
        contig_lengths[cid] = clen
        offset += clen
        full_seq_parts.append(seq)

    full_seq = "".join(full_seq_parts)
    full_seq = full_seq[rot:] + full_seq[:rot]
    gc_bins = []
    skew_bins = []
    for i in range(n_bins):
        st = int(i * total_len / n_bins)
        en = int((i + 1) * total_len / n_bins)
        sub = full_seq[st:en]
        g = sub.count("G")
        c = sub.count("C")
        atgc = max(1, len(sub))
        gc_bins.append(round((g + c) / atgc * 100.0, 2))
        skew_bins.append(round((g - c) / max(1, g + c), 4))

    mean_gc = round(sum(gc_bins) / len(gc_bins), 2)
    # predicted terminus: maximum of cumulative skew around the reordered genome
    cum, ter_bin, best = 0.0, None, None
    for i, v in enumerate(skew_bins):
        cum += v
        if best is None or cum > best:
            best, ter_bin = cum, i
    ter_bin = ter_bin if ori_pos is not None else None

    highlights = []
    for cid in contig_offsets:
        for r in rrna_by_contig[hab].get(cid, []):
            st = int(r["start"])
            en = int(r["end"])
            gpos = (contig_offsets[cid] + ((st + en) // 2)) % total_len
            rtype = r["rrna_type"].replace("_rRNA", " rRNA")
            highlights.append({
                "pos": gpos,
                "contig": cid,
                "contig_pos": st,
                "category": "rRNA",
                "label": rtype,
                "product": f"{rtype} ({r['gene_length']} bp, {float(r['gene_completeness'])*100:.0f}% complete)",
            })
        for r in trna_by_contig[hab].get(cid, []):
            st = int(r["start"])
            en = int(r["end"])
            gpos = (contig_offsets[cid] + ((st + en) // 2)) % total_len
            aa = r["amino_acid"]
            highlights.append({
                "pos": gpos,
                "contig": cid,
                "contig_pos": st,
                "category": "tRNA",
                "label": f"tRNA-{aa}",
                "product": f"tRNA-{aa} ({r['anticodon_or_tag']}, {r['gene_length']} bp)",
            })
        for r in bgc_by_contig[hab].get(cid, []):
            st = int(r["start"])
            en = int(r["end"])
            gpos = (contig_offsets[cid] + ((st + en) // 2)) % total_len
            btype = r["type"]
            sim = r["most_similar_known_cluster"]
            sim_str = f" (similar to {sim})" if sim and sim != "NA" else ""
            highlights.append({
                "pos": gpos,
                "contig": cid,
                "contig_pos": st,
                "category": "BGC",
                "label": f"BGC: {btype}",
                "product": f"antiSMASH BGC: {btype}{sim_str} [{st:,}-{en:,} bp]",
            })

    prots_by_contig = {}
    ko_desc_map = {}
    if os.path.exists(tsv_path):
        with open(tsv_path) as f:
            for r in csv.DictReader(f, delimiter="\t"):
                cid = r["contig_id"]
                if cid in contig_offsets:
                    prots_by_contig.setdefault(cid, []).append(r)
                ko = (r.get("KO") or "").replace("ko:", "").strip()
                desc = (r.get("description") or "").strip()
                if ko and desc and ko not in ko_desc_map:
                    ko_desc_map[ko] = desc
                for pf in (r.get("PFAMs") or "").split(","):
                    pf = pf.strip()
                    if pf and desc and pf not in ko_desc_map:
                        ko_desc_map[pf] = desc

    cds_bins = [0] * n_bins
    metabolic_bins = [0] * n_bins
    for cid, plist in prots_by_contig.items():
        n_p = max(1, len(plist))
        clen = contig_lengths[cid]
        coff = contig_offsets[cid]
        for idx, r in enumerate(plist):
            cpos = int((idx + 0.5) / n_p * clen)
            gpos = (coff + cpos) % total_len
            b_idx = min(n_bins - 1, max(0, int(gpos / total_len * n_bins)))
            cds_bins[b_idx] += 1
            ko = r.get("KO", "").replace("ko:", "")
            cazy = r.get("CAZy", "")
            pfam = r.get("PFAMs", "")
            desc = r.get("description", "")
            if ko or cazy:
                metabolic_bins[b_idx] += 1

            dlower = (desc + " " + pfam + " " + cazy + " " + ko).lower()
            if "sarcosine" in dlower:  # sarcosine oxidase subunits carry PFAM SoxA..SoxG, not sulfur oxidation
                dlower = dlower.replace("sox", "")
            cat = None
            lbl = None
            if cazy:
                cat = "CAZyme"
                lbl = f"CAZy: {cazy.split(',')[0]}"
            elif any(k in dlower for k in ["rubisco", "ribulose-bisphosphate carboxylase", "thiosulfate oxidation", "sox", "dsra", "dsrb", "adenylylsulfate reductase", "nitrite oxidoreductase", "rhodopsin", "bacteriorhodopsin", "cytochrome c oxidase"]):
                cat = "Energy"
                if "rubisco" in dlower or "ribulose-bisphosphate" in dlower: lbl = "RuBisCO (CBB)"
                elif "sox" in dlower or "thiosulfate" in dlower: lbl = "Sox (Sulfur Ox)"
                elif "dsr" in dlower or "sulfite" in dlower: lbl = "Dsr (Sulfur)"
                elif "nitrite oxidoreductase" in dlower: lbl = "Nxr (Nitrite Ox)"
                elif "rhodopsin" in dlower: lbl = "Rhodopsin (Solar)"
                else: lbl = "Cyt c Oxidase"
            elif any(k in dlower for k in ["sulfatase", "arylsulfatase", "alginate lyase", "tonb-dependent", "susc", "susd"]):
                cat = "CAZyme"
                if "sulfatase" in dlower: lbl = "Sulfatase"
                elif "alginate" in dlower: lbl = "Alginate Lyase"
                else: lbl = "TonB/SusC Receptor"
            elif any(k in dlower for k in ["biotin synthase", "thiamine", "dihydrofolate", "riboflavin synthase", "cobalamin", "glycine betaine", "osmoprotectant", "cold-shock", "cold shock", "siderophore"]):
                cat = "Vitamin/Cold"
                if "biotin" in dlower: lbl = "Biotin (Vit B7)"
                elif "folate" in dlower: lbl = "Folate (Vit B9)"
                elif "thiamine" in dlower: lbl = "Thiamine (Vit B1)"
                elif "riboflavin" in dlower: lbl = "Riboflavin (Vit B2)"
                elif "cobalamin" in dlower: lbl = "Cobalamin (Vit B12)"
                elif "cold" in dlower: lbl = "Cold-Shock (Csp)"
                elif "siderophore" in dlower: lbl = "Siderophore Transp."
                else: lbl = "Osmoprotectant"
            elif any(k in dlower for k in ["crispr", "cas1", "cas2", "cas3", "cas9"]):
                cat = "Defense"
                lbl = "CRISPR-Cas"

            if cat and lbl:
                highlights.append({
                    "pos": gpos,
                    "contig": cid,
                    "contig_pos": cpos,
                    "category": cat,
                    "label": lbl,
                    "product": " ".join(x for x in [desc.strip() or ko_desc_map.get(ko) or KO_NAMES.get(ko, ""), f"[{ko or pfam.split(',')[0]}]" if (ko or pfam) else ""] if x),
                })

    kept = []
    last_by_cat = {}
    priority_cats = {"rRNA": 0, "BGC": 0, "Defense": 1, "Energy": 1, "Vitamin/Cold": 2, "CAZyme": 2, "tRNA": 3}
    highlights.sort(key=lambda x: (priority_cats.get(x["category"], 9), x["pos"]))
    for h in highlights:
        cat = h["category"]
        if cat in ("rRNA", "BGC"):
            kept.append(h)
        else:
            min_dist = total_len * (0.025 if cat in ("Energy", "Defense") else 0.045)
            existing = last_by_cat.get(cat, [])
            if all(abs(h["pos"] - ep) >= min_dist for ep in existing) and len(existing) < 10:
                kept.append(h)
                existing.append(h["pos"])
                last_by_cat[cat] = existing

    kept.sort(key=lambda x: x["pos"])
    return {
        "total_bp": total_len,
        "n_contigs": len(contig_meta),
        "mean_gc": mean_gc,
        "contigs": contig_meta,
        "ori": {"method": ori_method, "contig": contigs[0][0] if ori_pos is not None else None},
        "ter_bin": ter_bin,
        "gc_bins": gc_bins,
        "skew_bins": skew_bins,
        "cds_bins": cds_bins,
        "metabolic_bins": metabolic_bins,
        "markers": kept,
    }, ko_desc_map

circos_all = {}
ko_desc_by_mag = {}
for m in mar_cands:
    circos_all[m], ko_desc_by_mag[m] = compute_circos("marine", m)
for m in fw_cands:
    circos_all[m], ko_desc_by_mag[m] = compute_circos("freshwater", m)

# ---------------------------------------------------------
# 3. ECOSSDB ECOSYSTEM SERVICES PER MAG
# ---------------------------------------------------------
cices_meta = {
    "2.1.1.1": {
        "short": "Bioremediation & Waste Decomposition",
        "cices_group": "Regulation & Maintenance · Mediation of Wastes & Toxic Substances",
        "sdgs": ["SDG 6: Clean Water", "SDG 14: Life Below Water", "SDG 12: Responsible Consumption"],
        "icon": "♻️",
        "desc": "Enzymatic breakdown and detoxification of organic detritus, hydrocarbons, and xenobiotic pollutants in cold waters."
    },
    "2.1.1.2": {
        "short": "Pollutant Filtration & Heavy-Metal Sequestration",
        "cices_group": "Regulation & Maintenance · Filtration & Sequestration by Biota",
        "sdgs": ["SDG 6: Clean Water", "SDG 3: Good Health", "SDG 14: Life Below Water"],
        "icon": "🛡️",
        "desc": "Active transport, chelation, and sequestration of heavy metals (arsenic, mercury, copper, iron) and toxic ions."
    },
    "2.3.5.1": {
        "short": "Controlling Chemical Quality of Freshwater",
        "cices_group": "Regulation & Maintenance · Water Conditions",
        "sdgs": ["SDG 6: Clean Water & Sanitation", "SDG 15: Life on Land"],
        "icon": "💧",
        "desc": "Cycling of nitrogen, phosphorus, sulfur, and dissolved organic carbon to keep tundra lakes clear and safe for fish and drinking."
    },
    "2.3.5.2": {
        "short": "Controlling Chemical Quality of Saltwater",
        "cices_group": "Regulation & Maintenance · Marine Water Conditions",
        "sdgs": ["SDG 14: Life Below Water", "SDG 6: Clean Water"],
        "icon": "🌊",
        "desc": "Marine sulfur oxidation, nutrient regeneration, and organic polymer turnover across tidal inlets and sea-ice interfaces."
    },
    "2.3.6.1": {
        "short": "Regulating Global Climate (Carbon & Trace Gases)",
        "cices_group": "Regulation & Maintenance · Atmospheric Composition & Conditions",
        "sdgs": ["SDG 13: Climate Action", "SDG 14: Life Below Water"],
        "icon": "🌍",
        "desc": "Fixation of CO₂, assimilation of C1 compounds (methanol/formate), and regulation of greenhouse gas fluxes."
    },
    "2.3.4.2": {
        "short": "Maintaining Sediment & Soil Organic Matter",
        "cices_group": "Regulation & Maintenance · Soil & Sediment Quality",
        "sdgs": ["SDG 15: Life on Land", "SDG 2: Zero Hunger"],
        "icon": "🌱",
        "desc": "Polysaccharide hydrolysis and humification that sustain benthic sediment fertility and food-web carbon."
    },
    "2.3.3.2": {
        "short": "Biological Disease & Pathogen Control",
        "cices_group": "Regulation & Maintenance · Pest & Disease Control",
        "sdgs": ["SDG 3: Good Health", "SDG 14: Life Below Water"],
        "icon": "🧬",
        "desc": "Antimicrobial secondary metabolites, quorum quenching, and competitive exclusion of aquatic pathogens."
    },
    "2.3.6.2": {
        "short": "Regulating Air Quality & Volatile Compounds",
        "cices_group": "Regulation & Maintenance · Atmospheric Quality",
        "sdgs": ["SDG 3: Good Health", "SDG 13: Climate Action"],
        "icon": "🌬️",
        "desc": "Transformation of volatile sulfur, nitrogen, and methylated compounds at the water-air interface."
    }
}

ecossdb_all = {}
es_files = [(hab, cands, f"{scratch}/{hab}/es_per_mag.tsv") for hab, cands in [("marine", mar_cands), ("freshwater", fw_cands)]]
es_files += [(v["habitat"], [g], f"{v['dir']}/es_per_mag.tsv") for g, v in EXTRA.items()]
for hab, cands, es_path in es_files:
    with open(es_path) as f:
        for r in csv.DictReader(f, delimiter="\t"):
            mid = r["entity_id"]
            if mid not in cands: continue
            code = r["es_code"]
            cm = cices_meta.get(code, {})
            # Parse roles e.g. "producer:461,transformer:2"
            roles_dict = {}
            for rp in (r.get("functional_roles") or "").split(","):
                if ":" in rp:
                    rk, rv = rp.split(":", 1)
                    try: roles_dict[rk.strip()] = int(rv.strip())
                    except: pass
            # Parse top_genes e.g. "K01560(0.80),K00249(0.70)"
            tgenes = []
            seen_g = set()
            for gp in (r.get("top_genes") or "").split(","):
                gp = gp.strip()
                if not gp: continue
                m_g = re.match(r"^([^\(]+)\(([0-9\.]+)\)$", gp)
                gid = m_g.group(1) if m_g else gp
                conf = float(m_g.group(2)) if m_g else 0.7
                if gid in seen_g: continue
                seen_g.add(gid)
                # KEGG's own KO name (short and complete) first, else the annotation's description
                gdesc = re.sub(r"\s*\[EC:[^\]]*\]$", "", KO_NAMES.get(gid, "")) or ko_desc_by_mag.get(mid, {}).get(gid, "")
                tgenes.append({
                    "id": gid,
                    "conf": conf,
                    "desc": (gdesc if len(gdesc) <= 80 else gdesc[:77].rstrip() + "…") if gdesc else gid,
                })
            ecossdb_all.setdefault(mid, []).append({
                "code": code,
                "name": r["es_name"],
                "short_name": cm.get("short", r["es_name"]),
                "cices_group": cm.get("cices_group", "CICES Regulation & Maintenance"),
                "sdgs": cm.get("sdgs", ["SDG 6: Clean Water"]),
                "icon": cm.get("icon", "🌿"),
                "desc": cm.get("desc", ""),
                "gene_count": int(r["gene_count"]),
                "weighted_score": round(float(r["weighted_score"]), 4),
                "roles": roles_dict,
                "top_genes": tgenes[:5],
            })

# ---------------------------------------------------------
# 4. NORTHERN GEOGRAPHIC RELATIVE ABUNDANCE MAP DATA
# ---------------------------------------------------------
# regional stations shown on the big map; is_ci marks the community's own station
station_coords = {st["name"]: {k: st.get(k) for k in ("lat", "lon", "region", "is_ci", "desc")}
                  for st in CFG["stations"]}

# The community's own sampling sites: coordinates, site and date per sample come from the
# field metadata TSV (the master sheet has no coordinates for these samples).
CI_META_TSV = cfg_path(LOCAL["metadata_tsv"])
CI_SITE_NAMES = LOCAL["site_names"]            # metadata `location` -> site shown on the page
CI_FRESH = set(LOCAL["freshwater_sites"])      # every other site counts as marine
ci_sample_meta = {}   # FLOWCELL_barcodeNN -> {site, loc_label, lat, lon, date}
ci_local_coords = {}  # site -> {lat, lon, name, type}
with open(CI_META_TSV) as f:
    for r in csv.DictReader(f, delimiter="\t"):
        loc = (r.get("location") or "").strip()
        site = CI_SITE_NAMES.get(loc)
        entry = {"site": site, "loc_label": loc.replace("_", " ") if site else "",
                 "station_id": r.get("station_id", ""), "date": (r.get("sample_date") or "").replace("-", "")}
        ci_sample_meta[f"{r['flowcell']}_{r['barcode']}"] = entry
        if site and r.get("lat") and r.get("lon"):
            ci_local_coords.setdefault(site, {"lat": float(r["lat"]), "lon": float(r["lon"]), "name": site,
                                              "type": "Freshwater" if site in CI_FRESH else "Marine"})

import openpyxl
SHEET = CFG["sample_sheet"]
wb = openpyxl.load_workbook(cfg_path(SHEET["xlsx"]), data_only=True)
ws = wb[SHEET["samples_tab"]]
meta = {}
for r in list(ws.iter_rows(values_only=True))[1:]:
    if not r or not r[0]: continue
    # the sheet writes barcode5, the depth tables barcode05
    fb = re.sub(r"_barcode(\d)$", r"_barcode0\1", str(r[0]).strip())
    meta[fb] = {
        "box": str(r[4] or "").strip(),
        "exp": str(r[5] or "").strip(),
        "sample_name": str(r[6] or "").strip(),
        "date": str(r[7] or "").strip(),
        "project": str(r[8] or "").strip(),
        "location": str(r[9] or "").strip(),
        "lat": r[12],
        "lon": r[13],
    }

# Comparison samples to leave out: runs marked failed on the "missing runs" sheet, and
# barcodes with under 25 Mb of reads. Local samples are always kept.
FAILED_RUNS = {str(r[1]).strip() for r in list(wb[SHEET["failed_runs_tab"]].iter_rows(values_only=True))[1:]
               if r and r[1] and str(r[4] or "").strip().lower() == "failed"}
barcode_mb = {}
for r in list(wb[SHEET["barcodes_tab"]].iter_rows(values_only=True))[1:]:
    try:
        barcode_mb[re.sub(r"_barcode(\d)$", r"_barcode0\1", str(r[0]).strip())] = float(r[10])
    except (TypeError, ValueError, IndexError):
        pass
MIN_MB = SHEET.get("min_mb", 25)
# freshwater regions: samples from these that ended up in the marine mapping are there by mistake
FRESHWATER_STATIONS = {st["name"] for st in CFG["stations"] if st.get("habitat") == "freshwater"}


def skip_sample(sn):
    return sn.split("_")[0] in FAILED_RUNS or barcode_mb.get(sn, MIN_MB) < MIN_MB


def habitat_hint(m):
    """Water type from the sample name/location: 'blank'; 'marine' for SW (seawater) and
    SF/SEAFOAM (sea foam); 'fresh' for FW, river and creek; None when the name doesn't say."""
    t = f"{m.get('sample_name', '')} {m.get('location', '')}".upper().replace(" ", "_")
    if re.search(r"BLANK", t):
        return "blank"
    if re.search(r"(^|[_-])(SW|SF)([_-]|$)|SEAFOAM", t):
        return "marine"
    if re.search(r"(^|[_-])FW([_-]|$)|RIVER|CREEK", t):
        return "fresh"
    return None


def classify_station(fb, m, hab):
    """First matching rule in build_config.json station_rules: `any` substrings of the
    upper-cased project/experiment/box/location/sample name (plus `and_regex` if given),
    or the project equal to `or_project`; else the habitat's fallback station."""
    proj = (m.get("project") or "").upper()
    exp  = (m.get("exp") or "").upper()
    box  = (m.get("box") or "").upper()
    loc  = (m.get("location") or "")
    sn   = (m.get("sample_name") or "")
    all_txt = f"{proj} {exp} {box} {loc} {sn}".upper()
    for r in CFG["station_rules"]:
        hit = any(w in all_txt for w in r.get("any", []))
        if hit and r.get("and_regex"):
            hit = re.search(r["and_regex"], all_txt) is not None
        if hit or (r.get("or_project") and proj == r["or_project"]):
            return r["station"]
    return CFG["fallback_station"][hab]

# Read-sketch hits from site_data_dump.txt cover local samples that have not been mapped yet
sketch_by_mag = {}
with open(f"{scratch}/site_data_dump.txt") as f:
    curr_m = None
    for line in f:
        if line.startswith("--- ") and line.strip().endswith(" ---"):
            curr_m = line.strip().strip("-").strip()
        elif line.startswith("Top sketch hits:") and curr_m:
            arr_str = line.split("Top sketch hits:", 1)[1].strip()
            sketch_by_mag[curr_m] = eval(arr_str)

for hits in sketch_by_mag.values():
    for h in hits:
        m = ci_sample_meta.get(h.get("sample"), {})
        h["loc"] = m.get("site") or ""
        h["loc_label"] = m.get("loc_label") or "not recorded"
        h["date"] = m.get("date") or h.get("date", "")

north_map_all = {}
MIN_BREADTH = CFG.get("min_breadth", 0.3)  # fallback rule: at least 30% of the genome has reads
# "detected" (depth-aware): at least `detect_min_bp` bases aligned to the genome at >=95% identity,
# and covered fraction at least `detect_breadth_ratio` of what that depth predicts for evenly
# spread reads (1 - e^-depth). A fixed breadth cut misses real but thinly sequenced genomes.
DETECT_MIN_BP = CFG.get("detect_min_bp", 250_000)
DETECT_RATIO = CFG.get("detect_breadth_ratio", 0.5)


def found(A, i, length):
    if not A["relab"][i] > 0:
        return False
    if not length:
        return A["breadth"][i] >= MIN_BREADTH
    d = A["depth"][i]
    return d * length >= DETECT_MIN_BP and A["breadth"][i] >= DETECT_RATIO * (1 - math.exp(-d))
unknown_stations = set()
jobs, hab_samples = [], {}
for hab, cands in [("marine", mar_cands), ("freshwater", fw_cands)]:
    # co-assembly MAGs: relative abundance from the co-assembly depth table (cand_abund.py on Grex),
    # unless `abundance_dir` + `mapped_bases` give every genome one measure (below)
    with open(f"{scratch}/{hab}/cand_abund.json") as f:
        ab = json.load(f)
    hab_samples[hab] = set(ab["samples"])
    idp = f"{scratch}/{hab}/cand_identity.json"
    ident = json.load(open(idp)) if os.path.exists(idp) else {}
    jobs.append((hab, cands, ab, ident, None))
# One measure for every genome (optional):
#   abundance_dir  <dir>/<id>/abund.json + identity.json from site_abund.sh, after all the site's
#                  genomes were mapped together (remap_genomes.sh): depth, breadth and identity
#                  from reads at >=95% identity
#   mapped_bases   {habitat: mapped_bases.tsv} from mapped_bases.sh on the co-assembly mapping:
#                  all bases mapped per sample
# relab = depth x length as % of ALL the bases that mapped to the habitat's co-assembly in that
# sample. (The co-assembly depths.txt is not used for this: its CoverM metabat-mode depths keep
# only a small part of long-read coverage.) Stand-alone genomes without these use their own
# folder's abund.json as is.
ABDIR = cfg_path(CFG["abundance_dir"]) if CFG.get("abundance_dir") else None
MAPPED = {}
for hab, path in (CFG.get("mapped_bases") or {}).items():
    MAPPED[hab] = {r[0]: int(r[2]) for r in csv.reader(open(cfg_path(path)), delimiter="\t") if len(r) > 2 and r[2].isdigit()}
remapped = {}
for g in mar_cands + fw_cands:
    hab = "marine" if g in mar_cands else "freshwater"
    d = f"{ABDIR}/{g}" if ABDIR and os.path.exists(f"{ABDIR}/{g}/abund.json") else (EXTRA[g]["dir"] if g in EXTRA else None)
    if not d:
        continue
    ab = json.load(open(f"{d}/abund.json"))
    a = ab["mags"][g]
    if hab in MAPPED:
        # a sample mapped only to the other habitat's co-assembly (a sample-sheet mix-up) uses that total
        other = MAPPED.get("freshwater" if hab == "marine" else "marine", {})
        tot = lambda s: MAPPED[hab].get(s) or (other.get(s) if s in ci_sample_meta else None)
        a["relab"] = [round(100 * dp * a["length"] / tot(s), 5) if tot(s) else 0
                      for s, dp in zip(ab["samples"], a["depth"])]
        allowed = set(MAPPED[hab]) | {s for s in other if s in ci_sample_meta}
    else:
        allowed = hab_samples[hab]
    idp = f"{d}/identity.json"
    remapped[g] = (hab, [g], ab, json.load(open(idp)) if os.path.exists(idp) else {}, allowed)
jobs = [(h, [c for c in cs if c not in remapped], a, i, al) for h, cs, a, i, al in jobs] + list(remapped.values())
for hab, cands, ab, ident, allowed in jobs:
    snames = ab["samples"]
    st_groups = {}
    for j, sn in enumerate(snames):
        if sn in LOCAL.get("exclude_samples", {}):
            continue  # e.g. an extraction control
        cm = ci_sample_meta.get(sn)
        if cm and cm.get("site"):
            if (cm["site"] in CI_FRESH) != (hab == "freshwater"):
                continue  # local sample from the other habitat
            st = LOCAL_STATION
        elif skip_sample(sn) or (allowed is not None and sn not in allowed):
            continue
        else:
            hint = habitat_hint(meta.get(sn, {}))
            if hint == "blank" or (hint == "marine" and hab == "freshwater") or (hint == "fresh" and hab == "marine"):
                continue  # control, or a sample mapped to the wrong habitat by mistake
            st = classify_station(sn, meta.get(sn, {}), hab)
            if hab == "marine" and st in FRESHWATER_STATIONS:
                continue
        if st not in station_coords:
            unknown_stations.add(st)
        st_groups.setdefault(st, []).append(j)

    for c in cands:
        A = ab["mags"].get(c, {"relab": [0] * len(snames), "depth": [0] * len(snames), "breadth": [0] * len(snames)})
        glen = A.get("length") or (dump_meta.get(c) or {}).get("size")
        sk_by_sample = {h["sample"]: h for h in sketch_by_mag.get(c, [])}

        regional_points = []
        for st, idxs in st_groups.items():
            sc = station_coords.get(st, {"lat": 55.0, "lon": -96.0, "region": st, "is_ci": False, "desc": ""})
            rel = [A["relab"][i] for i in idxs]
            regional_points.append({
                "station": st,
                "region": sc["region"],
                "lat": sc["lat"],
                "lon": sc["lon"],
                "is_ci": sc["is_ci"],
                "desc": sc["desc"],
                "n_samples": len(idxs),
                "detected": sum(1 for i in idxs if found(A, i, glen)),
                "mean_relab": round(sum(rel) / max(1, len(rel)), 4),
                "max_relab": round(max(rel, default=0.0), 4),
                "mean_cov": round(sum(A["depth"][i] for i in idxs) / max(1, len(idxs)), 4),
            })
        regional_points.sort(key=lambda x: (0 if x["is_ci"] else 1, -x["mean_relab"]))

        # Local samples of this habitat: mapped abundance where the sample is in the
        # depth table, otherwise the read-sketch match (samples not yet mapped)
        ci_rows = []
        for sn, cm in ci_sample_meta.items():
            if not cm.get("site") or (cm["site"] in CI_FRESH) != (hab == "freshwater"):
                continue
            sk = sk_by_sample.get(sn, {})
            row = {"id": cm.get("station_id") or sn, "sample": sn, "site": cm["site"], "loc_label": cm["loc_label"],
                   "date": cm.get("date", ""), "ani": round(sk["ani"], 2) if sk else None, "wkid": round(sk["wkid"], 1) if sk else None}
            idv = ident.get(sn, {}).get(c, {})
            if idv.get("identity") is not None and idv["reads"] < 5:
                row["ani"], row["ani_reads"], row["ani_source"] = None, idv["reads"], "few"  # 1-4 reads: too noisy
            elif idv.get("identity") is not None:
                row["ani"] = idv["identity"]  # mapped-read identity replaces the sketch estimate
                row["ani_reads"] = idv["reads"]
                row["ani_source"] = "mapped"
            elif sk:
                row["ani_source"] = "sketch"
            if sn in snames:
                i = snames.index(sn)
                row.update(source="mapped", relab=A["relab"][i], breadth=A["breadth"][i], depth=A["depth"][i],
                           detected=found(A, i, glen))
            elif sk:
                row.update(source="sketch", relab=None, breadth=None, depth=None, detected=True)
            else:
                continue
            ci_rows.append(row)
        ci_rows.sort(key=lambda r: (not r["detected"], -(r["relab"] if r["relab"] is not None else -1), -(r["wkid"] or 0)))

        local_by_loc = {}
        for r in ci_rows:
            if not r["detected"] or r["site"] not in ci_local_coords:
                continue
            e = local_by_loc.setdefault(r["site"], {**ci_local_coords[r["site"]], "loc_key": r["site"], "samples": [],
                                                    "max_relab": 0.0, "max_wkid": 0.0, "max_ani": 0.0})
            e["samples"].append(r)
            e["max_relab"] = max(e["max_relab"], r["relab"] or 0.0)
            e["max_wkid"] = max(e["max_wkid"], r["wkid"] or 0.0)
            e["max_ani"] = max(e["max_ani"], r["ani"] or 0.0)

        north_map_all[c] = {
            "habitat": hab,
            "regional_stations": regional_points,
            "ci_local_stations": list(local_by_loc.values()),
            "ci_samples": ci_rows,
        }
if unknown_stations:
    print(f"[WARN] stations without coordinates: {sorted(unknown_stations)}")

dump_meta, dump_evo, dump_mods = {}, {}, {}
with open(f"{scratch}/site_data_dump.txt") as f:
    curr_m = None
    for line in f:
        if line.startswith("--- ") and line.strip().endswith(" ---"):
            curr_m = line.strip().strip("-").strip()
        elif curr_m and line.startswith("MAG meta:"):
            dump_meta[curr_m] = eval(line.split(":", 1)[1].strip())
        elif curr_m and line.startswith("Evo:"):
            e = eval(line.split(":", 1)[1].strip())
            def g(*ks):
                for k in ks:
                    v = e.get(k)
                    if v not in (None, "N/A", ""):
                        return v
                return None
            dump_evo[curr_m] = {
                "classification": e.get("classification"),
                "method": g("method", "classification_method"),
                "closest_ref": g("closest_ref", "closest_genome_reference", "closest_placement_reference"),
                "closest_ani": g("closest_ani", "closest_genome_ani", "closest_placement_ani"),
                "closest_af": g("closest_af", "closest_genome_af", "closest_placement_af"),
                "red_value": g("red_value"),
                "msa_percent": g("msa_percent"),
                "other_refs": g("other_related_references(genome_id,species_name,radius,ANI,AF)"),
            }
        elif curr_m and line.startswith("Modules >=0.5:"):
            dump_mods[curr_m] = [{"id": k.split(":", 1)[0], "name": k.split(":", 1)[1], "c": v}
                                 for k, v in eval(line.split(":", 1)[1].strip())]

def evo_from_summary(r):
    """The evo fields the page shows, from one row of a GTDB-Tk summary."""
    def gv(*ks):
        for k in ks:
            if r.get(k) not in (None, "N/A", ""):
                return r[k]
        return None
    return {
        "classification": r["classification"],
        "method": gv("classification_method"),
        "closest_ref": gv("closest_genome_reference", "closest_placement_reference"),
        "closest_ani": gv("closest_genome_ani", "closest_placement_ani"),
        "closest_af": gv("closest_genome_af", "closest_placement_af"),
        "red_value": gv("red_value"),
        "msa_percent": gv("msa_percent"),
        "other_refs": gv("other_related_references(genome_id,species_name,radius,ANI,AF)"),
    }

# stand-alone genomes: quality from their own files, taxonomy from their GTDB-Tk summary
for g, v in EXTRA.items():
    dump_meta[g] = json.load(open(f"{v['dir']}/meta.json"))
    if v.get("domain") == "Archaea":
        dump_meta[g]["hypo_pct"] = None  # Bakta's database is built for bacteria: most archaeal genes come out "hypothetical"
    dump_mods[g] = [{"id": k.split(":", 1)[0], "name": k.split(":", 1)[1], "c": c}
                    for k, c in json.load(open(f"{v['dir']}/modules.json"))]
    for r in csv.DictReader(open(cfg_path(v["gtdbtk_summary"])), delimiter="\t"):
        if r["user_genome"] == g:
            dump_evo[g] = evo_from_summary(r)
for g, (r, tree) in GTDBTK_RUNS.items():  # newer GTDB-Tk runs override (see gtdbtk_runs)
    if g in dump_meta:
        dump_evo[g] = evo_from_summary(r)

# optional `canthyd`: canthyd_hits.tsv from scripts/hpc/canthyd.sh (CANT-HYD hydrocarbon-degradation
# marker genes per genome, with "confident"/"possible" from the HMMs' trusted and noise cutoffs)
canthyd_all = {}
if CFG.get("canthyd"):
    for r in csv.DictReader(open(cfg_path(CFG["canthyd"])), delimiter="\t"):
        canthyd_all.setdefault(r["genome"], []).append({
            "model": r["model"], "desc": r["description"], "score": float(r["score"]),
            "confidence": r["confidence"], "contig": r["contig"], "start": r["start"], "end": r["end"]})
    for v in canthyd_all.values():
        v.sort(key=lambda h: (h["confidence"] != "confident", h["model"]))

# closest relative and ANI from skani (ani_table), replacing GTDB-Tk's, which is blank when the
# nearest species is outside its ANI radius. No hit means nothing in GTDB is above ~80% ANI.
if ANI_HITS:
    for m, e in dump_evo.items():
        h = ANI_HITS.get(m, [])
        e.update({"ani_source": "skani", "ani_hits": h[:5],
                  "closest_ref": h[0]["acc"] if h else None, "closest_ani": h[0]["ani"] if h else None,
                  "closest_af": h[0]["af"] if h else None, "closest_species": h[0]["species"] if h else None})

# which steps of each pathway the genome has ("s": "1101", one character per step), for the
# step boxes on the Function tab. Steps come from site/explain/pathway_steps.js (danaSeq).
steps_js = os.path.join(os.path.dirname(os.path.abspath(__file__)), "../../site/explain/pathway_steps.js")
s = open(steps_js).read()
PATHWAY_STEPS = json.loads(s[s.index("= {") + 2:s.rstrip().rstrip(";").rindex("}") + 1])
for hab, ms in (("marine", mar_cands), ("freshwater", fw_cands)):
    for m in ms:
        kos = set()
        for r in csv.DictReader(open(f"{mag_dir(hab, m)}/{m}.tsv"), delimiter="\t"):
            kos.update(k.replace("ko:", "").strip() for k in (r.get("KO") or "").split(",") if k.strip())
        for mod in dump_mods.get(m, []):
            st = PATHWAY_STEPS.get(mod["id"])
            if st:
                mod["s"] = "".join("1" if kos & set(x["kos"]) else "0" for x in st["steps"])
                if abs(mod["s"].count("1") / len(mod["s"]) - mod["c"]) > 0.01:
                    print(f"[WARN] {m} {mod['id']}: steps {mod['s']} vs completeness {mod['c']}")

# ---------------------------------------------------------
# Home-page cladogram: each workshop MAG stands in for its closest GTDB reference
# genome, and the GTDB r226 bacterial reference tree is pruned to those genomes.
# Topology only (branch lengths dropped), so it shows how the MAGs relate.
# ---------------------------------------------------------
def refs_by_distance(tree):
    """GTDB reference leaves of an extracted tree dict, nearest to the target leaf first."""
    adj, leaves, target = {}, {}, None
    def walk(n, parent):
        k = id(n)
        adj.setdefault(k, [])
        if parent is not None:
            adj[k].append((parent, n.get("length", 0) or 0))
            adj[parent].append((k, n.get("length", 0) or 0))
        if n.get("children"):
            for ch in n["children"]:
                walk(ch, k)
        else:
            leaves[k] = n
    walk(tree["root"], None)
    target = next((k for k, n in leaves.items() if n.get("is_target")), None)
    if target is None:
        return []
    dist, stack = {target: 0.0}, [target]
    while stack:
        u = stack.pop()
        for v, l in adj[u]:
            if v not in dist:
                dist[v] = dist[u] + l
                stack.append(v)
    refs = sorted((dist[k], n["name"]) for k, n in leaves.items() if is_ref(n["name"]) and k in dist)
    return [r[1] for r in refs]

home_tree = None
_ref_tree = f"{scratch}/{CFG.get('reference_tree', 'bac120_r226.tree')}"
if os.path.exists(_ref_tree):
    gtdb_root = parse_newick(open(_ref_tree).read())
    in_tree = {lf.name.replace("GB_", "").replace("RS_", "") for lf in get_leaves(gtdb_root)}
    rep_of = {}
    for mag in mar_cands + fw_cands:
        # its closest reference, else the nearest reference leaf in its GTDB-Tk tree,
        # taking the first one that is in this release of the reference tree
        cands = [(dump_evo.get(mag) or {}).get("closest_ref")]
        if mag in extracted_trees:
            cands += refs_by_distance(extracted_trees[mag])
        for acc in cands:
            acc = (acc or "").replace("GB_", "").replace("RS_", "")
            if acc in in_tree and acc not in rep_of:
                rep_of[acc] = mag
                break
    keep = {}
    for lf in get_leaves(gtdb_root):
        acc = lf.name.replace("GB_", "").replace("RS_", "")
        if acc in rep_of:
            keep[id(lf)] = rep_of[acc]

    def prune(n):
        if not n.children:
            return {"id": keep[id(n)]} if id(n) in keep else None
        kids = [k for k in (prune(c) for c in n.children) if k]
        if not kids:
            return None
        return kids[0] if len(kids) == 1 else {"children": kids}
    home_tree = prune(gtdb_root)
    # archaea are not in the bacterial reference tree: a second branch at the root
    arch = [g for g in mar_cands + fw_cands if EXTRA.get(g, {}).get("domain") == "Archaea"]
    if arch:
        clade = {"id": arch[0]} if len(arch) == 1 else {"children": [{"id": g} for g in arch]}
        home_tree = {"children": [home_tree, clade]} if home_tree else clade
        keep.update({("arch", g): g for g in arch})
    missing = sorted(set(mar_cands + fw_cands) - set(keep.values()))
    print(f"Home tree: {len(keep)} MAGs placed via GTDB r226 references" + (f"; not placed: {missing}" if missing else ""))

bundle = {
    "homeTree": home_tree,
    "trees": extracted_trees,
    "circos": circos_all,
    "ecossdb": ecossdb_all,
    "northMap": north_map_all,
    "ciSites": list(ci_local_coords.values()),  # all local sampling sites (fixed local map extent)
    "meta": dump_meta,
    "evo": dump_evo,
    "modules": dump_mods,
    "sketch": sketch_by_mag,
    "canthyd": {m: canthyd_all.get(m, []) for m in mar_cands + fw_cands} if CFG.get("canthyd") else {},
    # cells per litre on the page = share of mapped DNA x total cells/L x (mean genome / this genome)
    "cellModel": {"cells_per_litre": CFG.get("cells_per_litre", {"marine": 1e9, "freshwater": 1e9}),
                  "mean_genome_bp": CFG.get("mean_genome_bp", 3_000_000)},
}

with open(out_js, "w") as f:
    f.write(f"// Auto-generated by build_site_data.py for {CFG.get('title', 'the naming workshop')}\n")
    f.write("window.EXTRA_SITE_DATA = ")
    json.dump(bundle, f, separators=(",", ":"))
    f.write(";\n")

print(f"Successfully wrote {out_js} ({os.path.getsize(out_js):,} bytes)")
